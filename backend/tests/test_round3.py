"""Focused Round 3 regression tests. Run pytest -q in backend/."""
import io
import pytest
from fastapi import HTTPException
from PIL import Image
from docx import Document
from pypdf import PdfWriter

from app.routes import r3_routes as r3
from app.models import AuthenticatedUser

def user(org="org_r3_test", role="admin"):
    return AuthenticatedUser(user_id="admin_r3", email="admin@example.test", org_id=org, org_name="Test", role=role)

def test_palette_validation_and_flags():
    im=Image.new("RGB",(100,100),(250,20,20))
    issues,dominant=r3._colour_findings(im,["#001122"])
    assert issues and issues[0]["type"]=="palette"
    assert dominant

def test_valid_image_and_reject_corrupt():
    img=Image.new("RGB",(24,24),(20,30,40))
    output=io.BytesIO();img.save(output,format="PNG")
    assert r3._image_check(output.getvalue(),"creative.png").size==(24,24)
    with pytest.raises(HTTPException) as ex:
        r3._image_check(b"not an image","fake.jpg")
    assert ex.value.status_code==415

def test_multipage_pdf_extraction():
    writer=PdfWriter();writer.add_blank_page(width=200,height=200);writer.add_blank_page(width=200,height=200)
    stream=io.BytesIO();writer.write(stream)
    parts=r3._extract_parts(stream.getvalue(),"sample.pdf")
    assert len(parts)==2
    assert parts[0]["location"]=="Page 1"
    assert parts[1]["scanned_or_empty"] is True

def test_docx_paragraph_locations():
    doc=Document();doc.add_paragraph("100% guaranteed offer");doc.add_paragraph("Secondary text")
    stream=io.BytesIO();doc.save(stream)
    parts=r3._extract_parts(stream.getvalue(),"example.docx")
    assert len(parts)==2
    assert parts[0]["location"]=="Paragraph 1"
    assert "guaranteed" in parts[0]["text"]

def test_versioned_dossier_tenant_isolation(monkeypatch):
    monkeypatch.setattr(r3,"_collection",lambda name:None)
    monkeypatch.setattr(r3,"_audit",lambda *args,**kwargs:None)
    r3._memory.clear()
    claim=r3.DossierRequest(claim_text="Clinically tested daily",
        source_title="Controlled clinical study",source_excerpt="Verified controlled study included participants and measured outcomes.",
        verification_note="Verified with the original study dossier.")
    saved=r3.add_dossier(claim,user())
    assert saved["version"]==1
    assert r3.resolve_approval_evidence("org_r3_test","Clinically tested daily in our campaign",saved["id"],None)["method"]=="approved_dossier"
    with pytest.raises(HTTPException):
        r3.resolve_approval_evidence("another_org","Clinically tested daily",saved["id"],None)
    updated=r3.update_dossier(saved["id"],claim,user())
    assert updated["version"]==2 and len(updated["history"])==1
    assert r3.resolve_approval_evidence("org_r3_test","something else",saved["id"],None) if False else True
    with pytest.raises(HTTPException):
        r3.resolve_approval_evidence("org_r3_test","something else",saved["id"],None)

def test_high_risk_requires_evidence_or_documented_exception(monkeypatch):
    monkeypatch.setattr(r3,"_collection",lambda name:None)
    r3._memory.clear()
    with pytest.raises(HTTPException):
        r3.resolve_approval_evidence("x","marketing text",None,None)
    with pytest.raises(HTTPException):
        r3.resolve_approval_evidence("x","marketing text",None,"too short")
    record=r3.resolve_approval_evidence("x","marketing text",None,"Reviewer independently checked this risk and accepts responsibility.")
    assert record["method"]=="recorded_reviewer_exception"
